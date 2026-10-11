"""Make the fake case folder for the dress run: briefs/brief-a.md split into BRIEF.md + a docx + an xlsx + README.md + SSS.md + api/openapi.json.   python3 scripts/make_dress_case.py <out folder>"""
import re, sys
from pathlib import Path
import docx, openpyxl
T = Path(sys.argv[1]); T.mkdir(parents=True, exist_ok=True)
secs = [s for s in re.split(r"(?m)^(?=#{1,3} )", Path("briefs/brief-a.md").read_text(encoding="utf-8")) if s.strip()]
n = len(secs)
(T / "BRIEF.md").write_text(secs[0] + "".join(secs[1:max(2, n // 2)]))
d = docx.Document(); d.add_heading("Duyuru / Ek kurallar", 1)
for s in secs[max(2, n // 2):n - 1]:
    for line in s.splitlines():
        if line.strip():
            d.add_paragraph(line.lstrip("# ").strip())
d.save(str(T / "ek_kurallar.docx"))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Tablo"; ws.append(["Gün", "Kural"]); ws.append(["Pazartesi", "İlk arama 09:00"])
for line in secs[-1].splitlines()[:8]:
    if line.strip():
        ws.append(["Not", line.strip()[:200]])
wb.save(str(T / "tablo.xlsx"))
(T / "README.md").write_text("# Vaka README\nBu klasör tek bir vakadır. Türkçe notlar: ş ğ ü ı İ.\n")
(T / "SSS.md").write_text("# SSS\nSoru: Kural değişti mi? Cevap: Ek kurallar dosyasına bakın.\n")
(T / "api").mkdir(exist_ok=True); (T / "api" / "openapi.json").write_text(Path("tools/openapi.json").read_text())
print("dress case in", T)
