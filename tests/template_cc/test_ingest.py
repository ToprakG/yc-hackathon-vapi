import json
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))


def test_ingest_order_headers_and_api_fallback(tmp_path, monkeypatch):
    import docx, openpyxl
    from builder import ingest as I
    c = tmp_path / "case"; c.mkdir()
    (c / "BRIEF.md").write_text("ana kural şğü"); (c / "README.md").write_text("readme")
    (c / "aaa_sss.md").write_text("soru cevap")
    d = docx.Document(); d.add_paragraph("duyuru metni"); t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "a"; t.rows[0].cells[1].text = "b"; d.save(str(c / "duyuru.docx"))
    wb = openpyxl.Workbook(); wb.active.append(["x", "y"]); wb.active.append([1, 2]); wb.save(str(c / "t.xlsx"))
    (c / "spec.json").write_text(json.dumps({"openapi": "3.0.0", "paths": {"/a": {}}, "servers": [{"url": "https://real.example/v1"}]}))
    monkeypatch.setattr(I, "fetch_openapi", lambda u: (_ for _ in ()).throw(OSError("down")))
    r = I.ingest(c, "http://mock.example", tmp_path / "out")
    text = Path(r["brief"]).read_text()
    heads = [l for l in text.splitlines() if l.startswith("=== FILE:")]
    assert heads[:2] == ["=== FILE: BRIEF.md ===", "=== FILE: README.md ==="] and len(heads) == 5
    assert "duyuru metni" in text and "| x | y |" in text and "ana kural şğü" in text and "spec.json" not in text
    assert json.loads(Path(r["api"]).read_text())["servers"][0]["url"] == "https://real.example/v1"
