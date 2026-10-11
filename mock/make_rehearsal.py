"""Rehearsal only (nothing here goes into the builder): a Turkish case folder that mirrors the final format, made from brief-a.

    python3 mock/make_rehearsal.py     -> mock/rehearsal-tr/FreyaBank-Test/ (BRIEF.md, README.md, duyuru.docx, servis-sozlesmesi.pdf, musteriler.xlsx, SSS.md, ornek-gorusme-1/2.md)
The PDF needs fpdf2 (pip install --target /tmp/pdflib fpdf2; PYTHONPATH=/tmp/pdflib) and the Arial font of macOS."""
import importlib.util
import json
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "mock"))
import convert as C  # noqa: E402

OUT = REPO / "mock/rehearsal-tr/FreyaBank-Test"
OUT.mkdir(parents=True, exist_ok=True)
brief = (REPO / "briefs/brief-a.md").read_text()
openapi = (REPO / "tools/openapi.json").read_text()
spec = json.loads((REPO / "mock/spec.brief-a.json").read_text())

if not (OUT / "BRIEF.md").exists():
    tr = C.llm("You are a precise translator.", "Translate this bank call-handling brief into Turkish. Keep the structure, every rule number, every number, date format, "
               "tool name, field name and enum value exactly as they are (tool and field names stay in English). Output only the translated Markdown.\n\n" + brief, 16000)
    (OUT / "BRIEF.md").write_text(tr)
tr_brief = (OUT / "BRIEF.md").read_text()

PROMPT = """Write rehearsal documents for a Turkish bank phone-agent case. Output ONLY a JSON object with these string fields, all in Turkish:
- readme: a short README.md (what each file is: BRIEF.md, duyuru.docx, servis-sozlesmesi.pdf, musteriler.xlsx, SSS.md, ornek-gorusme-1.md, ornek-gorusme-2.md)
- announcement: the text of an internal announcement (docx) that CHANGES exactly one numeric or date rule of the brief (pick one, e.g. a limit or a time window; say clearly it replaces the old value, quoting the old value) - 2 short paragraphs
- contract: the service contract in Turkish prose: every endpoint of the OpenAPI below (path, method, request fields, response fields, error behaviour, that identity must pass first, that the service does not enforce the brief) as headed sections
- qa: a SSS.md with 6 questions and answers an agent builder might ask about the brief (answers must agree with the brief, except the one the announcement changes)
- conv1, conv2: two example conversations as Markdown (AGENT:/MÜŞTERİ: lines, then the tool calls made and the closing outcome); conv1 = a normal promise; conv2 = a customer who asks 'bu bir dolandırıcılık mı?' and then refuses. Use only facts, numbers and tool names from the brief and the OpenAPI.
BRIEF (Turkish):
%s

OPENAPI:
%s"""
docs_path = OUT.parent / ".docs.json"
if not docs_path.exists():
    docs_path.write_text(json.dumps(C.json_out(C.llm("You write precise documents.", PROMPT % (tr_brief, openapi), 20000)), ensure_ascii=False))
d = json.loads(docs_path.read_text())
(OUT / "README.md").write_text(d["readme"])
(OUT / "SSS.md").write_text(d["qa"])
(OUT / "ornek-gorusme-1.md").write_text(d["conv1"])
(OUT / "ornek-gorusme-2.md").write_text(d["conv2"])

if "--pdf" not in sys.argv:
    import docx  # noqa: E402
    doc = docx.Document()
    doc.add_heading("Duyuru", 1)
    for para in d["announcement"].split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    doc.save(str(OUT / "duyuru.docx"))

    # xlsx of the test customers (the Turkish export of tests/mock/test_convert.py)
    spec_t = importlib.util.spec_from_file_location("tc", REPO / "tests/mock/test_convert.py")
    tc = importlib.util.module_from_spec(spec_t)
    spec_t.loader.exec_module(tc)
    tc._turkish_xlsx(OUT / "musteriler.xlsx", spec)


# pdf: made by `python3 mock/make_rehearsal.py --pdf` (system python 3.11+ with fpdf2; the repo venv is 3.10)
if "--pdf" not in sys.argv:
    import subprocess
    subprocess.run(["python3", __file__, "--pdf"], check=True, env={**__import__("os").environ, "PYTHONPATH": "/tmp/pdflib"})
else:
    # pdf: the contract prose + the openapi JSON as an appendix
    sys.path.insert(0, "/tmp/pdflib")
    from fpdf import FPDF  # noqa: E402
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("Arial", "", "/System/Library/Fonts/Supplemental/Arial.ttf")
    pdf.add_font("Arial", "B", "/System/Library/Fonts/Supplemental/Arial Bold.ttf")
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 8, "Servis Sözleşmesi", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", 10)
    for line in d["contract"].splitlines():
        pdf.multi_cell(0, 5, line.lstrip("# ").strip() or " ", new_x="LMARGIN", new_y="NEXT")
    pdf.add_page()
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 8, "Ek: OpenAPI (JSON)", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", 7)
    for line in json.dumps(json.loads(openapi), indent=1, ensure_ascii=False).splitlines():
        pdf.multi_cell(0, 3.2, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(OUT / "servis-sozlesmesi.pdf"))
    print("wrote", sorted(p.name for p in OUT.iterdir()))
